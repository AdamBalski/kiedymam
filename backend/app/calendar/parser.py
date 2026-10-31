from __future__ import annotations

from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook


class ScheduleParseError(ValueError):
    pass


WEEKDAYS = {
    "PONIEDZIALEK": "Monday",
    "WTOREK": "Tuesday",
    "SRODA": "Wednesday",
    "CZWARTEK": "Thursday",
    "PIATEK": "Friday",
    "SOBOTA": "Saturday",
    "NIEDZIELA": "Sunday",
}
TIME_RANGE = re.compile(r"^\s*(\d{1,2})[.:](\d{2})\s*[-–]\s*(\d{1,2})[.:](\d{2})\s*$")
LEADING_TIME_RANGE = re.compile(r"^\s*\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}\s*(?:\([WSC]\)\s*)?", re.IGNORECASE)
GROUP_PREFIX = re.compile(
    r"^\s*(?P<label>S\s*\d{1,3}|\d{1,3}(?:\s*\.\s*)?\s*\(\s*\d{1,3}\s*\)|\d{1,3}(?:[.,]\d{1,3})+|\d{1,3})"
    r"(?P<separator>\.)?\s+(?P<body>.+)$",
    re.IGNORECASE,
)
GROUP_CAPACITY = re.compile(r"\((\d{1,2})\)")
CHANGE_DATE = re.compile(r"^\s*(\d{1,2})[.:](\d{1,2})\s*:\s*$")


def _fold(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(c for c in normalized if not unicodedata.combining(c))


def _weekday(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    letters = "".join(c for c in _fold(value).upper() if c.isalpha())
    for polish, english in WEEKDAYS.items():
        if letters.startswith(polish):
            return english
    return None


def _time_range(value: object) -> tuple[time, time] | None:
    if not isinstance(value, str):
        return None
    match = TIME_RANGE.match(value)
    if not match:
        return None
    h1, m1, h2, m2 = (int(part) for part in match.groups())
    if h1 > 23 or h2 > 24 or m1 > 59 or m2 > 59 or (h2 == 24 and m2 != 0):
        return None
    return time(h1, m1), time(0, 0) if (h2, m2) == (24, 0) else time(h2, m2)


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _group_info(text: str) -> tuple[str | None, list[str], str]:
    match = GROUP_PREFIX.match(text)
    if not match:
        return None, [], text.strip()
    label = re.sub(r"\s+", "", match.group("label")).upper()
    if label.startswith("S"):
        numbers = [label[1:]]
    else:
        numbers = re.findall(r"\d+", label.split("(", 1)[0])
    capacity_match = GROUP_CAPACITY.search(label)
    if capacity_match:
        numbers.append(f"size:{capacity_match.group(1)}")
    return label, numbers, match.group("body").strip()


def _clean_summary(text: str) -> str:
    return LEADING_TIME_RANGE.sub("", text).strip() or text.strip()


def _base_group_numbers(value: object) -> list[int]:
    if isinstance(value, bool) or value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    if re.fullmatch(r"\d{1,3}(?:\.\d{1,3})+", text):
        parts = [int(part) for part in text.split(".")]
        for index in range(1, len(parts)):
            previous, current = parts[index - 1], parts[index]
            # Excel stores a displayed pair such as "19.20" as numeric 19.2.
            if previous >= 10 and previous % 10 == 9 and current == 2:
                parts[index] = previous + 1
        return parts
    if re.fullmatch(r"\d{1,3}", text):
        number = int(text)
        return [number] if 1 <= number <= 999 else []
    return []


def _group_rows(sheet, header_row: int, section_end: int) -> dict[int, int]:
    rows: dict[int, int] = {}
    for row_number in range(header_row + 1, min(section_end, header_row + 8) + 1):
        value = sheet.cell(row_number, 1).value
        try:
            group_size = int(value)
            if group_size in (6, 12, 18, 36):
                rows[group_size] = row_number
        except (TypeError, ValueError):
            continue
    return rows


def _groups_from_header(value: object, group_size: int) -> set[int]:
    if group_size == 6:
        return set(_base_group_numbers(value))
    text = str(value).strip().upper() if value is not None else ""
    if group_size == 36:
        match = re.match(r"^S\s*(\d+)$", text)
        group_number = int(match.group(1)) if match else None
    else:
        match = re.match(r"^(\d+)\s*\.?\s*\(\s*(\d+)\s*\)$", text)
        group_number = int(match.group(1)) if match and int(match.group(2)) == group_size else None
    if group_number is None:
        return set()
    first = (group_number - 1) * (group_size // 6) + 1
    return set(range(first, first + group_size // 6))


def _group_header_ranges(sheet, group_rows: dict[int, int], merged_bounds: dict) -> dict[int, list[tuple[int, int, object]]]:
    header_ranges: dict[int, list[tuple[int, int, object]]] = {}
    for group_size, row_number in group_rows.items():
        ranges = []
        for col in range(2, sheet.max_column + 1):
            cell = sheet.cell(row_number, col)
            if cell.value is None:
                continue
            _, header_max_col = merged_bounds.get((row_number, col), (row_number, col))
            ranges.append((col, header_max_col, cell.value))
        header_ranges[group_size] = ranges
    return header_ranges


def _groups_for_span(header_ranges: dict[int, list[tuple[int, int, object]]], min_col: int, max_col: int) -> list[int]:
    selected: set[int] = set()
    for col in range(min_col, max_col + 1):
        memberships = []
        for group_size, ranges in header_ranges.items():
            label_set = next((
                _groups_from_header(value, group_size)
                for start, end, value in ranges
                if start <= col <= end
            ), set())
            if label_set:
                memberships.append(label_set)
        if memberships:
            selected.update(set.intersection(*memberships))
    return sorted(group for group in selected if 1 <= group <= 999)


def _groups_from_event_label(label: str | None) -> list[int]:
    if not label:
        return []
    text = label.strip().upper()
    if text.startswith("S"):
        return sorted(_groups_from_header(text, 36))
    match = re.fullmatch(r"(\d{1,3})\s*\.?\s*\(\s*(\d{1,3})\s*\)", text)
    if match:
        number, size = int(match.group(1)), int(match.group(2))
        if size == 6:
            return [number]
        return sorted(_groups_from_header(f"{number}({size})", size))
    return sorted({int(value) for value in re.findall(r"\d+", text) if 1 <= int(value) <= 999})


def _changelog(workbook) -> list[dict]:
    entries: list[dict] = []
    for sheet in workbook.worksheets:
        comment = sheet["B2"].comment
        if not comment or not comment.text.strip():
            continue
        current: dict | None = None
        for line in comment.text.splitlines():
            text = line.strip()
            match = CHANGE_DATE.match(text)
            if match:
                if current:
                    entries.append(current)
                current = {"date": f"{int(match.group(1)):02d}.{int(match.group(2)):02d}", "items": []}
            elif current and text:
                current["items"].append(text.lstrip("-• ").strip())
        if current:
            entries.append(current)
    return entries


def _date_headers(sheet, row_number: int, max_column: int | None = None) -> list[tuple[int, date]]:
    headers: list[tuple[int, date]] = []
    for col in range(2, (max_column or sheet.max_column) + 1):
        parsed = _as_date(sheet.cell(row_number, col).value)
        if parsed:
            headers.append((col, parsed))
    return headers


def _date_for_column(headers: list[tuple[int, date]], col: int) -> date | None:
    selected = None
    for header_col, header_date in headers:
        if header_col > col:
            break
        selected = header_date
    return selected


def _week_block_for_column(headers: list[tuple[int, date]], col: int, max_column: int) -> tuple[int, int] | None:
    selected_index = None
    for index, (header_col, _) in enumerate(headers):
        if header_col > col:
            break
        selected_index = index
    if selected_index is None:
        return None
    start = headers[selected_index][0]
    if selected_index + 1 < len(headers):
        end = headers[selected_index + 1][0] - 1
    elif len(headers) > 1:
        end = start + (headers[-1][0] - headers[-2][0]) - 1
    else:
        end = max_column
    return start, min(end, max_column)


def _merged_bounds(sheet) -> dict[tuple[int, int], tuple[int, int]]:
    bounds: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in sheet.merged_cells.ranges:
        bounds[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)
    return bounds


def parse_g0_workbook(contents: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(contents), data_only=True, read_only=False)
    except Exception as exc:
        raise ScheduleParseError("The XLSX document could not be opened") from exc

    parsed: list[dict] = []
    available_groups: set[int] = set()
    for sheet in workbook.worksheets:
        day_headers: list[tuple[int, str, list[tuple[int, date]]]] = []
        for row_number in range(1, sheet.max_row + 1):
            weekday = _weekday(sheet.cell(row_number, 1).value)
            if weekday:
                dates = _date_headers(sheet, row_number)
                if dates:
                    day_headers.append((row_number, weekday, dates))
        if not day_headers:
            continue

        merged_bounds = _merged_bounds(sheet)
        for day_index, (header_row, weekday, dates) in enumerate(day_headers):
            section_end = day_headers[day_index + 1][0] - 1 if day_index + 1 < len(day_headers) else sheet.max_row
            group_rows = _group_rows(sheet, header_row, section_end)
            group_header_ranges = _group_header_ranges(sheet, group_rows, merged_bounds)
            for group_size, ranges in group_header_ranges.items():
                for _start, _end, label in ranges:
                    available_groups.update(_groups_from_header(label, group_size))
            for row_number in range(header_row + 1, section_end + 1):
                interval = _time_range(sheet.cell(row_number, 1).value)
                if not interval:
                    continue
                starts_at, _ = interval
                for col in range(2, sheet.max_column + 1):
                    cell = sheet.cell(row_number, col)
                    if not isinstance(cell.value, str) or not cell.value.strip():
                        continue
                    event_date = _date_for_column(dates, col)
                    if not event_date:
                        continue
                    week_block = _week_block_for_column(dates, col, sheet.max_column)
                    if not week_block:
                        continue
                    week_start_col, week_end_col = week_block
                    end_row, end_col = merged_bounds.get((row_number, col), (row_number, col))
                    end_interval = _time_range(sheet.cell(end_row, 1).value)
                    if not end_interval:
                        continue
                    _, ends_at = end_interval
                    raw = " ".join(cell.value.replace("\n", " ").split())
                    group_label, group_numbers, summary = _group_info(raw)
                    summary = _clean_summary(summary)
                    coordinate_groups = _groups_for_span(group_header_ranges, col, end_col)
                    all_groups = col <= week_start_col and end_col >= week_end_col
                    event = {
                        "sheet": sheet.title,
                        "weekday": weekday,
                        "date": event_date.isoformat(),
                        "start": starts_at.strftime("%H:%M"),
                        "end": ends_at.strftime("%H:%M"),
                        "summary": summary or raw,
                        "source_text": raw,
                        "group_label": group_label,
                        "group_numbers": [] if all_groups else (coordinate_groups or _groups_from_event_label(group_label)),
                        "all_groups": all_groups,
                        "source_cell": cell.coordinate,
                        "source_range": f"{cell.coordinate}:{sheet.cell(end_row, end_col).coordinate}",
                    }
                    event["event_key"] = sha256(
                        "|".join(
                            str(event[key])
                            for key in ("sheet", "date", "group_label", "summary")
                        ).encode("utf-8")
                    ).hexdigest()
                    parsed.append(event)

    if not parsed:
        workbook.close()
        raise ScheduleParseError("No dated schedule entries were found in this workbook")
    unique = {event["event_key"]: event for event in parsed}
    result = {
        "events": sorted(unique.values(), key=lambda event: (event["date"], event["start"], event["sheet"], event["source_cell"])),
        "changelog": _changelog(workbook),
        "_available_groups": sorted(available_groups),
    }
    workbook.close()
    return result


def parse_workbook(contents: bytes, parser_key: str) -> dict:
    from app.calendar.parsers.registry import parse_workbook as parse_registered_workbook

    return parse_registered_workbook(contents, parser_key)


def parse_workbook_with_format(contents: bytes, format_id: str) -> dict:
    from app.calendar.parsers.registry import parse_workbook_with_format as parse_registered_format

    return parse_registered_format(contents, format_id)
