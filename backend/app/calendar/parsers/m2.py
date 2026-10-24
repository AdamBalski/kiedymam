from __future__ import annotations

from datetime import date
from hashlib import sha256
from io import BytesIO
import re

from openpyxl import load_workbook

from app.calendar.parser import ScheduleParseError, _changelog, _fold, _time_range

DATE_IN_TEXT = re.compile(
    r"(?<!\d)(?:(?P<iso_year>20\d{2})[-/.](?P<iso_month>\d{1,2})[-/.](?P<iso_day>\d{1,2})|"
    r"(?P<day>\d{1,2})\.(?P<month>\d{1,2})(?:\.(?P<year>\d{2,4}))?)(?!\d)"
)
TIME_IN_TEXT = re.compile(r"(?<!\d)(\d{1,2}[.:]\d{2})\s*[-–]\s*(\d{1,2}[.:]\d{2})(?!\d)")
ACADEMIC_YEAR = re.compile(r"(20\d{2})\s*[/\-]\s*(20\d{2})")
DATE_RANGE_ONLY = re.compile(
    r"\s*\d{1,2}[./]\d{1,2}(?:[./]\d{2,4})?\s*[-–]\s*"
    r"\d{1,2}[./]\d{1,2}(?:[./]\d{2,4})?\s*"
)
GROUP_LIST = re.compile(
    r"(?:GRUP(?:A|Y)?\s*[:.]?\s*)"
    r"([1-9]\d{0,2}(?:\s*[,;. ]\s*[1-9]\d{0,2})*(?:\s*[-–]\s*[1-9]\d{0,2})?)",
    re.IGNORECASE,
)
WEEKDAYS = {"poniedzialek", "wtorek", "sroda", "czwartek", "piatek", "sobota", "niedziela"}


def _academic_year(workbook) -> tuple[int, int]:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    match = ACADEMIC_YEAR.search(cell.value)
                    if match:
                        return int(match.group(1)), int(match.group(2))
    return date.today().year, date.today().year + 1


def _dates_in_text(value: object, academic_year: tuple[int, int]) -> list[date]:
    if not isinstance(value, str):
        return []
    result: list[date] = []
    for match in DATE_IN_TEXT.finditer(value):
        if match.group("iso_year"):
            year, month, day = int(match.group("iso_year")), int(match.group("iso_month")), int(match.group("iso_day"))
        else:
            day, month = int(match.group("day")), int(match.group("month"))
            raw_year = match.group("year")
            if raw_year:
                year = int(raw_year)
                if year < 100:
                    year += 2000
            else:
                year = academic_year[0] if month >= 9 else academic_year[1]
        try:
            parsed = date(year, month, day)
        except ValueError:
            continue
        if parsed not in result:
            result.append(parsed)
    return result


def _time_from_text(value: object):
    if not isinstance(value, str):
        return None
    match = TIME_IN_TEXT.search(value)
    if not match:
        return None
    return _time_range(f"{match.group(1)}-{match.group(2)}")


def _numbers(value: object) -> list[int]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    match = GROUP_LIST.search(text)
    if match:
        numeric_text = match.group(1)
    elif re.fullmatch(r"\s*[1-9]\d{0,2}(?:\s*[,;. ]\s*[1-9]\d{0,2})*(?:\s*[-–]\s*[1-9]\d{0,2})?\s*[\"']?\s*", text):
        numeric_text = text
    else:
        return []
    numbers = [int(part) for part in re.findall(r"\d+", numeric_text)]
    if "-" in numeric_text or "–" in numeric_text:
        start, end = numbers[:2]
        numbers = list(range(start, end + 1)) if start <= end else []
    return sorted({number for number in numbers if 1 <= number <= 999})


def _event(sheet, cell, event_date: date, interval, summary: str, source_text: str, groups: list[int]) -> dict:
    start, end = interval
    if hasattr(event_date, "date"):
        event_date = event_date.date()
    event = {
        "sheet": sheet.title,
        "weekday": (
            "Monday",
            "Tuesday",
            "Wednesday",
            "Thursday",
            "Friday",
            "Saturday",
            "Sunday",
        )[event_date.weekday()],
        "date": event_date.isoformat(),
        "start": start.strftime("%H:%M"),
        "end": end.strftime("%H:%M"),
        "summary": summary.strip() or source_text.strip(),
        "source_text": source_text.strip(),
        "group_label": None,
        "group_numbers": groups,
        "all_groups": not groups,
        "source_cell": cell.coordinate,
        "source_range": cell.coordinate,
    }
    event["event_key"] = sha256(
        "|".join(
            str(event[key])
            for key in ("sheet", "date", "start", "end", "group_label", "group_numbers", "summary")
        ).encode("utf-8")
    ).hexdigest()
    return event


def _header_columns(sheet, row_number: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row_number, col).value
        if not isinstance(value, str):
            continue
        label = " ".join(_fold(value).casefold().split())
        if label.startswith("przedmiot"):
            result["subject"] = col
        elif label.startswith("dzien"):
            result["day"] = col
        elif label.startswith("godzina"):
            result["time"] = col
        elif label.startswith("uwagi") or label.startswith("termin"):
            result["notes"] = col
        elif label.startswith("sala") or label.startswith("miejsce"):
            result["room"] = col
    return result


def _table_events(sheet, academic_year: tuple[int, int]) -> tuple[list[dict], set[str]]:
    events: list[dict] = []
    handled: set[str] = set()
    for row_number in range(1, sheet.max_row + 1):
        columns = _header_columns(sheet, row_number)
        if "subject" not in columns or "time" not in columns:
            continue
        subject_col, time_col = columns["subject"], columns["time"]
        notes_col = columns.get("notes")
        room_col = columns.get("room")
        current_subject = ""
        for data_row in range(row_number + 1, sheet.max_row + 1):
            subject_value = sheet.cell(data_row, subject_col).value
            if isinstance(subject_value, str) and subject_value.strip():
                if DATE_RANGE_ONLY.fullmatch(subject_value):
                    current_subject = ""
                    continue
                folded = _fold(subject_value).strip().casefold()
                if folded in {"wyklady", "seminaria", "cwiczenia", "cwiczenia praktyczne"}:
                    current_subject = ""
                    continue
                current_subject = " ".join(subject_value.split())
            if not current_subject:
                continue
            time_value = sheet.cell(data_row, time_col).value
            interval = _time_from_text(time_value)
            notes_value = sheet.cell(data_row, notes_col).value if notes_col else None
            dates = _dates_in_text(notes_value, academic_year)
            if not interval or not dates:
                continue
            room_value = sheet.cell(data_row, room_col).value if room_col else None
            source_text = " | ".join(
                str(value).strip()
                for value in (current_subject, room_value, notes_value)
                if value
            )
            groups = _numbers(current_subject)
            for event_date in dates:
                event = _event(
                    sheet,
                    sheet.cell(data_row, subject_col),
                    event_date,
                    interval,
                    current_subject,
                    source_text,
                    groups,
                )
                events.append(event)
            handled.add(sheet.cell(data_row, notes_col).coordinate if notes_col else "")
    return events, handled


def _date_grid_events(sheet, academic_year: tuple[int, int]) -> list[dict]:
    events: list[dict] = []
    date_rows: list[tuple[int, list[tuple[int, date]]]] = []
    for row_number in range(1, sheet.max_row + 1):
        headers = []
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row_number, col).value
            if isinstance(value, (date,)):
                headers.append((col, value))
        if len(headers) >= 2:
            date_rows.append((row_number, headers))
    for index, (header_row, headers) in enumerate(date_rows):
        section_end = date_rows[index + 1][0] - 1 if index + 1 < len(date_rows) else sheet.max_row
        first_date_col = headers[0][0]
        group_col = max(1, first_date_col - 1)
        for row_number in range(header_row + 1, section_end + 1):
            label = sheet.cell(row_number, group_col).value
            if isinstance(label, str) and _fold(label).strip().casefold() in WEEKDAYS | {"grupa"}:
                continue
            groups = _numbers(label)
            for col in range(first_date_col, sheet.max_column + 1):
                cell = sheet.cell(row_number, col)
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                interval = _time_from_text(cell.value)
                if not interval:
                    continue
                event_date = None
                for date_col, candidate in headers:
                    if date_col > col:
                        break
                    event_date = candidate
                if event_date is None:
                    continue
                dates = _dates_in_text(TIME_IN_TEXT.sub("", cell.value), academic_year)
                selected_dates = dates or [event_date]
                summary = TIME_IN_TEXT.sub("", cell.value).strip()
                if not summary:
                    continue
                for item in selected_dates:
                    events.append(_event(sheet, cell, item, interval, summary, cell.value, groups))
    return events


def _free_text_events(sheet, academic_year: tuple[int, int], handled: set[str]) -> list[dict]:
    events: list[dict] = []
    for row in sheet.iter_rows():
        group_candidates = [_numbers(cell.value) for cell in row[: min(4, len(row))]]
        row_groups = max(group_candidates, key=len, default=[])
        for cell in row:
            if cell.coordinate in handled or not isinstance(cell.value, str):
                continue
            interval = _time_from_text(cell.value)
            dates = _dates_in_text(TIME_IN_TEXT.sub("", cell.value), academic_year)
            if not interval or not dates:
                continue
            summary = TIME_IN_TEXT.sub("", cell.value)
            summary = DATE_IN_TEXT.sub("", summary)
            summary = re.sub(r"(?i)w\s*termin(?:ach|ie)\s*:", "", summary)
            if not summary.strip():
                continue
            for event_date in dates:
                events.append(_event(sheet, cell, event_date, interval, summary, cell.value, row_groups))
    return events


def parse_m2_workbook(contents: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(contents), data_only=True, read_only=False)
    except Exception as exc:
        raise ScheduleParseError("The XLSX document could not be opened") from exc

    year = _academic_year(workbook)
    available_groups = {
        group
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        for group in _numbers(cell.value)
    }
    events: list[dict] = []
    for sheet in workbook.worksheets:
        table_events, handled = _table_events(sheet, year)
        events.extend(table_events)
        events.extend(_date_grid_events(sheet, year))
        events.extend(_free_text_events(sheet, year, handled))

    if not events:
        workbook.close()
        raise ScheduleParseError("No dated Medicine session rows were found in this workbook")

    unique = {event["event_key"]: event for event in events}
    result = {
        "events": sorted(
            unique.values(),
            key=lambda event: (
                event["date"], event["start"], event["sheet"], event["source_cell"]
            ),
        ),
        "changelog": _changelog(workbook),
        "_available_groups": sorted(available_groups),
    }
    workbook.close()
    return result


class M2Parser:
    format_id = "M2"

    def parse(self, contents: bytes) -> dict:
        return parse_m2_workbook(contents)
