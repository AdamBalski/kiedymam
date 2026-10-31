from __future__ import annotations

from datetime import date
from hashlib import sha256
from io import BytesIO
import re

from openpyxl import load_workbook

from app.calendar.parser import (
    ScheduleParseError,
    _changelog,
    _clean_summary,
    _date_for_column,
    _date_headers,
    _fold,
    _groups_from_event_label,
    _groups_from_header,
    _group_info,
    _merged_bounds,
    _time_range,
    _week_block_for_column,
    _weekday,
)

GROUP_SIZE = re.compile(
    r"(?:GRUP(?:A|Y)?\s*)?(6|8|12|18|24|36)\s*(?:OS(?:OBY|OBOWE)?)?",
    re.IGNORECASE,
)
NUMBERED_GROUP = re.compile(r"^\s*(\d{1,3})\s*[\"']?\s*$")
TEXT_DATE = re.compile(r"(?<!\d)(\d{1,2})[./](\d{1,2})[./](20\d{2})(?!\d)")
LEADING_SLOT_TIME = re.compile(r"^\s*\d{1,2}[.:]\d{2}\s*[-–]\s*")
WEEKDAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
ROMAN_VALUES = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
    "VI": 6,
    "VII": 7,
    "VIII": 8,
    "IX": 9,
    "X": 10,
    "XI": 11,
    "XII": 12,
    "XIII": 13,
    "XIV": 14,
    "XV": 15,
    "XVI": 16,
    "XVII": 17,
    "XVIII": 18,
    "XIX": 19,
    "XX": 20,
    "XXI": 21,
    "XXII": 22,
    "XXIII": 23,
    "XXIV": 24,
    "XXV": 25,
    "XXVI": 26,
    "XXVII": 27,
    "XXVIII": 28,
    "XXIX": 29,
    "XXX": 30,
    "XXXI": 31,
    "XXXII": 32,
    "XXXIII": 33,
    "XXXIV": 34,
    "XXXV": 35,
    "XXXVI": 36,
}


def _row_group_size(value: object) -> int | None:
    if not isinstance(value, str):
        return None
    match = GROUP_SIZE.search(_fold(value).upper())
    return int(match.group(1)) if match else None


def _date_headers_for_format(
    sheet, row_number: int, format_id: str, max_column: int
) -> list[tuple[int, date]]:
    headers = _date_headers(sheet, row_number, max_column)
    if format_id != "M1":
        return headers

    for col in range(1, max_column + 1):
        value = sheet.cell(row_number, col).value
        if not isinstance(value, str) or not _weekday(value):
            continue
        match = TEXT_DATE.search(value)
        if not match:
            continue
        day, month, year = (int(part) for part in match.groups())
        try:
            headers.append((col, date(year, month, day)))
        except ValueError:
            continue
    return sorted({col: value for col, value in headers}.items())


def _matrix_date_rows(sheet, format_id: str) -> list[tuple[int, list[tuple[int, date]]]]:
    candidates: list[tuple[int, list[tuple[int, date]]]] = []
    max_row = sheet.max_row
    max_column = sheet.max_column
    for row_number in range(1, max_row + 1):
        headers = _date_headers_for_format(sheet, row_number, format_id, max_column)
        if not headers:
            continue
        preceding_weekday = (
            any(
                _weekday(sheet.cell(row_number - 1, col).value)
                for col in range(1, min(max_column, 12) + 1)
            )
            if row_number > 1
            else False
        )
        current_weekday = any(
            _weekday(sheet.cell(row_number, col).value)
            for col in range(1, max_column + 1)
        )
        has_termin_marker = any(
            "termin" in _fold(str(sheet.cell(row_number, col).value or "")).casefold()
            for col in range(1, max_column + 1)
        )
        group_heading = any(
            _row_group_size(sheet.cell(candidate, col).value) is not None
            or "grupa" in _fold(str(sheet.cell(candidate, col).value or "")).casefold()
            for candidate in range(row_number + 1, min(max_row, row_number + 7) + 1)
            for col in range(1, max_column + 1)
        )
        if format_id == "M1":
            if (preceding_weekday or current_weekday) and group_heading:
                candidates.append((row_number, headers))
        elif group_heading and (preceding_weekday or has_termin_marker):
            candidates.append((row_number, headers))
    return candidates


def _choose_time_column(sheet, header_row: int, section_end: int, max_column: int) -> int:
    candidate_columns = range(1, max_column + 1)
    return max(
        candidate_columns,
        key=lambda col: sum(
            _time_range(sheet.cell(row, col).value) is not None
            for row in range(header_row + 1, section_end + 1)
        ),
    )


def _group_header_rows(
    sheet,
    header_row: int,
    first_time_row: int,
    format_id: str,
    group_label_col: int,
    first_date_col: int,
    max_column: int,
) -> list[tuple[int, int]]:
    rows: list[tuple[int, int]] = []
    upper_bound = min(first_time_row - 1, header_row + 8)
    for row_number in range(header_row + 1, upper_bound + 1):
        label = sheet.cell(row_number, group_label_col).value
        group_size = _row_group_size(label)
        if group_size is not None:
            rows.append((row_number, group_size))
        elif format_id == "M1" and isinstance(label, str) and "grupa" in _fold(label).casefold():
            rows.append((row_number, 6))
        elif format_id == "M1":
            labels_in_row = [
                sheet.cell(row_number, col).value
                for col in range(first_date_col, max_column + 1)
            ]
            has_group_heading = any(
                isinstance(value, str) and "grupa" in _fold(value).casefold()
                for value in labels_in_row
            )
            if has_group_heading:
                rows.append((row_number + 1, 6))
    return rows


def _header_ranges(
    sheet,
    rows: list[tuple[int, int]],
    bounds: dict,
    first_date_col: int,
    max_col: int,
) -> dict[int, list[tuple[int, int, object]]]:
    result: dict[int, list[tuple[int, int, object]]] = {}
    for row_number, group_size in rows:
        values: list[tuple[int, int, object]] = []
        for col in range(first_date_col, max_col + 1):
            value = sheet.cell(row_number, col).value
            if value is None:
                continue
            _, end_col = bounds.get((row_number, col), (row_number, col))
            values.append((col, end_col, value))
        result[group_size] = values
    return result


def _groups_in_span(
    ranges: dict[int, list[tuple[int, int, object]]],
    min_col: int,
    max_col: int,
) -> list[int]:
    selected: set[int] = set()
    for col in range(min_col, max_col + 1):
        memberships: list[set[int]] = []
        for group_size, group_ranges in ranges.items():
            header = next((
                value for start, end, value in group_ranges if start <= col <= end
            ), None)
            if header is None:
                continue
            ids = _groups_from_header(header, group_size)
            if not ids:
                match = NUMBERED_GROUP.fullmatch(str(header))
                if match:
                    ids = [int(match.group(1))]
                else:
                    roman = _fold(str(header)).strip().upper()
                    if roman in ROMAN_VALUES:
                        ids = [ROMAN_VALUES[roman]]
            if ids:
                memberships.append(set(ids))
        if memberships:
            selected.update(set.intersection(*memberships))
    return sorted(value for value in selected if 1 <= value <= 999)


def _weekday_for_date(value: date) -> str:
    return WEEKDAY_NAMES[value.weekday()]


def _make_event(
    sheet,
    cell,
    event_date: date,
    start,
    end,
    all_groups: bool,
    groups: list[int],
    raw: str,
) -> dict:
    cleaned_raw = LEADING_SLOT_TIME.sub("", raw)
    group_label, _, summary = _group_info(cleaned_raw)
    summary = _clean_summary(summary)
    event_groups = groups or _groups_from_event_label(group_label)
    event = {
        "sheet": sheet.title,
        "weekday": _weekday_for_date(event_date),
        "date": event_date.isoformat(),
        "start": start.strftime("%H:%M"),
        "end": end.strftime("%H:%M"),
        "summary": summary or cleaned_raw,
        "source_text": raw,
        "group_label": group_label,
        "group_numbers": [] if all_groups else event_groups,
        "all_groups": all_groups,
        "source_cell": cell.coordinate,
        "source_range": cell.coordinate,
    }
    event["event_key"] = sha256(
        "|".join(
            str(event[key])
            for key in ("sheet", "date", "start", "end", "group_label", "group_numbers", "summary")
        ).encode("utf-8")
    ).hexdigest()
    return event


def parse_matrix_workbook(contents: bytes, format_id: str) -> dict:
    try:
        workbook = load_workbook(BytesIO(contents), data_only=True, read_only=False)
    except Exception as exc:
        raise ScheduleParseError("The XLSX document could not be opened") from exc

    events: list[dict] = []
    available_groups: set[int] = set()
    for sheet in workbook.worksheets:
        sheet_max_row = sheet.max_row
        sheet_max_column = sheet.max_column
        date_rows = _matrix_date_rows(sheet, format_id)
        if not date_rows:
            continue
        bounds = _merged_bounds(sheet)
        for header_index, (header_row, headers) in enumerate(date_rows):
            section_end = (
                date_rows[header_index + 1][0] - 1
                if header_index + 1 < len(date_rows)
                else sheet_max_row
            )
            first_date_col = headers[0][0]
            time_col = _choose_time_column(sheet, header_row, section_end, sheet_max_column)
            time_rows = [
                row for row in range(header_row + 1, section_end + 1)
                if _time_range(sheet.cell(row, time_col).value)
            ]
            if not time_rows:
                continue
            last_date_col = headers[-1][0]
            if len(headers) > 1:
                date_steps = sorted(
                    right[0] - left[0]
                    for left, right in zip(headers, headers[1:])
                )
                date_step = max(1, date_steps[len(date_steps) // 2])
            else:
                date_step = 1
            grid_last_col = min(sheet_max_column, last_date_col + date_step - 1)
            if time_col > last_date_col:
                grid_last_col = min(grid_last_col, time_col - 1)
            group_label_col = (
                time_col if time_col > last_date_col else max(1, first_date_col - 1)
            )
            group_rows = _group_header_rows(
                sheet,
                header_row,
                time_rows[0],
                format_id,
                group_label_col,
                first_date_col,
                sheet_max_column,
            )
            group_ranges = _header_ranges(sheet, group_rows, bounds, first_date_col, grid_last_col)
            for group_size, ranges in group_ranges.items():
                for _start, _end, label in ranges:
                    available_groups.update(_groups_from_header(label, group_size))
            for row_number in time_rows:
                start, _ = _time_range(sheet.cell(row_number, time_col).value)
                for col in range(first_date_col, grid_last_col + 1):
                    if col == time_col:
                        continue
                    cell = sheet.cell(row_number, col)
                    if not isinstance(cell.value, str) or not cell.value.strip():
                        continue
                    event_date = _date_for_column(headers, col)
                    block = _week_block_for_column(headers, col, grid_last_col)
                    if event_date is None or block is None:
                        continue
                    end_row, end_col = bounds.get((row_number, col), (row_number, col))
                    interval = _time_range(sheet.cell(end_row, time_col).value)
                    if interval is None:
                        continue
                    _, end = interval
                    raw = " ".join(cell.value.replace("\n", " ").split())
                    if _time_range(raw):
                        continue
                    groups = _groups_in_span(group_ranges, col, end_col)
                    all_groups = col <= block[0] and end_col >= block[1]
                    if format_id == "M1" and groups == list(range(1, 7)):
                        all_groups = True
                    events.append(_make_event(sheet, cell, event_date, start, end, all_groups, groups, raw))

    if not events:
        workbook.close()
        raise ScheduleParseError(f"No dated schedule entries were found in this {format_id} workbook")

    unique = {event["event_key"]: event for event in events}
    result = {
        "events": sorted(
            unique.values(),
            key=lambda event: (
                event["date"], event["start"], event["sheet"], event["source_cell"]
            ),
        ),
        "changelog": _changelog(workbook),
        "_available_groups": sorted(available_groups),
    }
    workbook.close()
    return result
